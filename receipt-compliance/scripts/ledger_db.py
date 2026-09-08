#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
进销项台账与月度统计引擎（v4.4.0）

把散落的票据识别结果沉淀为可分析的资产：
1. SQLite 台账主表（含认证期限字段）
2. 报表视图：月度汇总、税负率趋势、供应商集中度、异常波动（z-score）
3. Excel 导出（openpyxl，不可用时自动降级 CSV）

设计约束：
- 数据库文件默认落在 ~/.workbuddy/output/，不污染 skill 仓库
- openpyxl 非标准库，缺失时降级为 CSV，功能不中断
"""

import json
import sqlite3
import statistics
import sys
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional, Dict, Any, List, Iterable

try:
    from unified_invoice import UnifiedInvoice
except ImportError:  # 允许从其他目录调用
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from unified_invoice import UnifiedInvoice


DEFAULT_DB_PATH = Path.home() / ".workbuddy" / "output" / "receipt_ledger.db"

# 认证期限默认窗口（天）
# 说明：自 2020-03-01 起，增值税专用发票已取消 360 天认证确认期限。
# 此处保留窗口仅作为企业内部「勾选所属期」管理提醒，可在 YAML 中调整或关闭。
DEFAULT_CERT_WINDOW_DAYS = 360

DIRECTIONS = ("input", "output")


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _month_of(date_str: Optional[str]) -> Optional[str]:
    """从日期串提取 YYYY-MM"""
    if not date_str:
        return None
    s = date_str.strip().replace("/", "-")
    if len(s) >= 7:
        return s[:7]
    return None


def _to_float(v: Any) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


class LedgerDB:
    """进销项台账（SQLite）"""

    SCHEMA = """
    CREATE TABLE IF NOT EXISTS invoices (
        id                     INTEGER PRIMARY KEY AUTOINCREMENT,
        direction              TEXT NOT NULL DEFAULT 'input',
        invoice_type           TEXT,
        receipt_type           TEXT,
        invoice_code           TEXT,
        invoice_number         TEXT,
        billing_date           TEXT,
        month                  TEXT,
        amount                 REAL DEFAULT 0,
        tax_amount             REAL DEFAULT 0,
        total                  REAL DEFAULT 0,
        seller_name            TEXT,
        seller_tax_id          TEXT,
        buyer_name             TEXT,
        buyer_tax_id           TEXT,
        expense_category       TEXT,
        vat_deduction_amount   REAL DEFAULT 0,
        certification_deadline TEXT,
        certification_source   TEXT,
        source_file            TEXT,
        raw_json               TEXT,
        created_at             TEXT,
        UNIQUE(invoice_code, invoice_number, direction)
    );
    CREATE INDEX IF NOT EXISTS idx_inv_month     ON invoices(month);
    CREATE INDEX IF NOT EXISTS idx_inv_seller    ON invoices(seller_name);
    CREATE INDEX IF NOT EXISTS idx_inv_direction ON invoices(direction);
    CREATE INDEX IF NOT EXISTS idx_inv_deadline  ON invoices(certification_deadline);
    """

    def __init__(self, db_path: Optional[str] = None):
        self.db_path = Path(db_path) if db_path else DEFAULT_DB_PATH
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(self.db_path))
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(self.SCHEMA)
        self.conn.commit()

    # ---------- 写入 ----------

    def upsert_invoice(self, inv: UnifiedInvoice, direction: str = "input",
                       source_file: Optional[str] = None) -> int:
        """写入或更新一张票据，返回主键 id"""
        if direction not in DIRECTIONS:
            raise ValueError(f"direction 必须是 {DIRECTIONS} 之一，收到 {direction}")

        billing_date = inv.billing_date or inv.travel_date or ""
        row = {
            "direction": direction,
            "invoice_type": inv.invoice_type,
            "receipt_type": inv.receipt_type,
            "invoice_code": inv.invoice_code,
            "invoice_number": inv.invoice_number,
            "billing_date": billing_date,
            "month": _month_of(billing_date),
            "amount": _to_float(inv.amount),
            "tax_amount": _to_float(inv.tax_amount),
            "total": _to_float(inv.total),
            "seller_name": inv.seller_name,
            "seller_tax_id": inv.seller_tax_id,
            "buyer_name": inv.buyer_name,
            "buyer_tax_id": inv.buyer_tax_id,
            "expense_category": inv.expense_category,
            "vat_deduction_amount": _to_float(inv.vat_deduction_amount),
            "source_file": source_file,
            "raw_json": inv.to_json(),
            "created_at": _now(),
        }

        cur = self.conn.execute(
            "SELECT id, certification_deadline, certification_source FROM invoices "
            "WHERE invoice_code IS ? AND invoice_number IS ? AND direction = ?",
            (row["invoice_code"], row["invoice_number"], direction),
        )
        exist = cur.fetchone()

        if exist:
            # 已存在：只更新业务字段，保留人工补录的认证期限
            cols = [c for c in row if c != "created_at"]
            sets = ", ".join(f"{c} = ?" for c in cols)
            self.conn.execute(
                f"UPDATE invoices SET {sets} WHERE id = ?",
                [row[c] for c in cols] + [exist["id"]],
            )
            self.conn.commit()
            return exist["id"]

        # 新增：全电发票自动推算认证期限
        if inv.invoice_type == "full_electronic" and billing_date:
            row["certification_deadline"] = self.compute_deadline(billing_date)
            row["certification_source"] = "auto"

        cols = list(row.keys())
        sql = (f"INSERT INTO invoices ({', '.join(cols)}) "
               f"VALUES ({', '.join('?' * len(cols))})")
        cur = self.conn.execute(sql, [row[c] for c in cols])
        self.conn.commit()
        return cur.lastrowid

    def upsert_many(self, items: Iterable, direction: str = "input") -> List[int]:
        """批量写入，items 为 UnifiedInvoice 或 (invoice, direction, source_file)"""
        ids = []
        for item in items:
            if isinstance(item, tuple):
                inv, d, src = (list(item) + [None, None])[:3]
                ids.append(self.upsert_invoice(inv, d or direction, src))
            else:
                ids.append(self.upsert_invoice(item, direction))
        return ids

    def load_from_json(self, json_path: str, direction: str = "input") -> int:
        """从票据解析结果 JSON 批量导入"""
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            data = data.get("invoices") or data.get("results") or []
        count = 0
        for item in data:
            known = {k: v for k, v in item.items() if hasattr(UnifiedInvoice(), k)}
            inv = UnifiedInvoice(**known)
            self.upsert_invoice(inv, direction, item.get("source_file"))
            count += 1
        return count

    # ---------- 认证期限 ----------

    @staticmethod
    def compute_deadline(billing_date: str, window_days: int = DEFAULT_CERT_WINDOW_DAYS) -> Optional[str]:
        """按开票日期推算认证期限"""
        if not billing_date:
            return None
        s = billing_date.strip().replace("/", "-")
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y年%m月%d日"):
            try:
                d = datetime.strptime(s, fmt)
                return (d + timedelta(days=window_days)).strftime("%Y-%m-%d")
            except ValueError:
                continue
        return None

    def set_certification_deadline(self, invoice_number: str, deadline: str,
                                   invoice_code: Optional[str] = None) -> bool:
        """
        手工补录认证期限（纸票）
        不传 invoice_code 时只按发票号码匹配
        """
        if invoice_code:
            cur = self.conn.execute(
                "UPDATE invoices SET certification_deadline = ?, certification_source = 'manual' "
                "WHERE invoice_number = ? AND invoice_code = ?",
                (deadline, invoice_number, invoice_code),
            )
        else:
            cur = self.conn.execute(
                "UPDATE invoices SET certification_deadline = ?, certification_source = 'manual' "
                "WHERE invoice_number = ?",
                (deadline, invoice_number),
            )
        self.conn.commit()
        return cur.rowcount > 0

    def fill_missing_deadlines(self, window_days: int = DEFAULT_CERT_WINDOW_DAYS) -> int:
        """为缺失认证期限的记录按开票日期补算"""
        cur = self.conn.execute(
            "SELECT id, billing_date FROM invoices "
            "WHERE (certification_deadline IS NULL OR certification_deadline = '') "
            "  AND billing_date IS NOT NULL AND billing_date != ''"
        )
        rows = cur.fetchall()
        n = 0
        for r in rows:
            dl = self.compute_deadline(r["billing_date"], window_days)
            if dl:
                self.conn.execute(
                    "UPDATE invoices SET certification_deadline = ?, certification_source = 'auto' "
                    "WHERE id = ?", (dl, r["id"]))
                n += 1
        self.conn.commit()
        return n

    # ---------- 查询 ----------

    def query_invoices(self, direction: Optional[str] = None,
                       start: Optional[str] = None,
                       end: Optional[str] = None) -> List[Dict[str, Any]]:
        sql = "SELECT * FROM invoices WHERE 1=1"
        args: List[Any] = []
        if direction:
            sql += " AND direction = ?"
            args.append(direction)
        if start:
            sql += " AND billing_date >= ?"
            args.append(start)
        if end:
            sql += " AND billing_date <= ?"
            args.append(end)
        sql += " ORDER BY billing_date, id"
        return [dict(r) for r in self.conn.execute(sql, args).fetchall()]

    # ---------- 报表视图 ----------

    def monthly_summary(self) -> List[Dict[str, Any]]:
        """月度进销项汇总"""
        rows = self.conn.execute(
            """
            SELECT month,
                   direction,
                   COUNT(*)                    AS count,
                   ROUND(SUM(amount), 2)       AS amount,
                   ROUND(SUM(tax_amount), 2)   AS tax_amount,
                   ROUND(SUM(total), 2)        AS total
            FROM invoices
            WHERE month IS NOT NULL AND month != ''
            GROUP BY month, direction
            ORDER BY month
            """
        ).fetchall()
        return [dict(r) for r in rows]

    def tax_burden_trend(self, months: int = 12) -> List[Dict[str, Any]]:
        """
        税负率趋势
        税负率 = (销项税额 - 进项税额) / 不含税销售额
        """
        rows = self.conn.execute(
            """
            SELECT month,
                   SUM(CASE WHEN direction='output' THEN amount     ELSE 0 END) AS out_amount,
                   SUM(CASE WHEN direction='output' THEN tax_amount  ELSE 0 END) AS out_tax,
                   SUM(CASE WHEN direction='input'  THEN tax_amount  ELSE 0 END) AS in_tax,
                   SUM(CASE WHEN direction='input'  THEN amount      ELSE 0 END) AS in_amount
            FROM invoices
            WHERE month IS NOT NULL AND month != ''
            GROUP BY month
            ORDER BY month DESC
            LIMIT ?
            """, (months,)
        ).fetchall()

        result = []
        for r in reversed([dict(x) for x in rows]):
            out_amount = r["out_amount"] or 0
            payable = (r["out_tax"] or 0) - (r["in_tax"] or 0)
            burden = (payable / out_amount) if out_amount else 0.0
            result.append({
                "month": r["month"],
                "output_amount": round(out_amount, 2),
                "output_tax": round(r["out_tax"] or 0, 2),
                "input_tax": round(r["in_tax"] or 0, 2),
                "payable_tax": round(payable, 2),
                "tax_burden": round(burden, 4),
                "tax_burden_pct": f"{burden * 100:.2f}%",
            })
        return result

    def supplier_concentration(self, top_n: int = 5) -> Dict[str, Any]:
        """
        供应商集中度：Top N 供应商占进项总额比例及环比变化
        """
        rows = self.conn.execute(
            """
            SELECT seller_name, month,
                   SUM(total) AS total
            FROM invoices
            WHERE direction = 'input' AND seller_name IS NOT NULL AND seller_name != ''
            GROUP BY seller_name, month
            """
        ).fetchall()

        by_supplier: Dict[str, float] = {}
        by_month: Dict[str, Dict[str, float]] = {}
        grand_total = 0.0
        for r in rows:
            name = r["seller_name"]
            val = r["total"] or 0
            by_supplier[name] = by_supplier.get(name, 0) + val
            grand_total += val
            if r["month"]:
                by_month.setdefault(r["month"], {})
                by_month[r["month"]][name] = by_month[r["month"]].get(name, 0) + val

        ordered = sorted(by_supplier.items(), key=lambda kv: kv[1], reverse=True)
        top = ordered[:top_n]
        top_sum = sum(v for _, v in top)
        top_share = (top_sum / grand_total) if grand_total else 0.0

        # 最近两个月的 Top1 占比变化
        months = sorted(by_month.keys())
        change = None
        if len(months) >= 2:
            prev_m, last_m = months[-2], months[-1]
            prev_total = sum(by_month[prev_m].values())
            last_total = sum(by_month[last_m].values())
            prev_top1 = max(by_month[prev_m].values()) if by_month[prev_m] else 0
            last_top1 = max(by_month[last_m].values()) if by_month[last_m] else 0
            prev_share = (prev_top1 / prev_total) if prev_total else 0
            last_share = (last_top1 / last_total) if last_total else 0
            change = {
                "from_month": prev_m,
                "to_month": last_m,
                "top1_share_prev": round(prev_share, 4),
                "top1_share_last": round(last_share, 4),
                "delta": round(last_share - prev_share, 4),
            }

        return {
            "grand_total": round(grand_total, 2),
            "top_n": top_n,
            "top_suppliers": [
                {"supplier": n, "total": round(v, 2),
                 "share": round(v / grand_total, 4) if grand_total else 0.0}
                for n, v in top
            ],
            "top_share": round(top_share, 4),
            "top_share_pct": f"{top_share * 100:.2f}%",
            "concentration_change": change,
        }

    def amount_anomaly(self, z_threshold: float = 2.5) -> Dict[str, Any]:
        """
        异常波动：对月度金额序列做 z-score，超过阈值即标注
        """
        rows = self.conn.execute(
            """
            SELECT month, SUM(total) AS total, COUNT(*) AS count
            FROM invoices
            WHERE month IS NOT NULL AND month != ''
            GROUP BY month ORDER BY month
            """
        ).fetchall()
        series = [{"month": r["month"], "total": r["total"] or 0, "count": r["count"]}
                  for r in rows]

        if len(series) < 3:
            return {"months": len(series), "mean": 0, "std": 0,
                    "z_threshold": z_threshold, "flagged": [],
                    "note": "样本月数不足 3，z-score 不具统计意义"}

        values = [s["total"] for s in series]
        mean = statistics.mean(values)
        std = statistics.pstdev(values) or 0.0

        flagged = []
        for s in series:
            z = (s["total"] - mean) / std if std else 0.0
            s["z_score"] = round(z, 3)
            if abs(z) >= z_threshold:
                flagged.append({**s, "direction": "偏高" if z > 0 else "偏低"})

        return {
            "months": len(series),
            "mean": round(mean, 2),
            "std": round(std, 2),
            "z_threshold": z_threshold,
            "flagged": flagged,
            "series": series,
        }

    def full_report(self, months: int = 12, top_n: int = 5,
                    z_threshold: float = 2.5) -> Dict[str, Any]:
        """一次产出全部报表视图"""
        return {
            "report_date": _now(),
            "db_path": str(self.db_path),
            "invoice_count": self.conn.execute("SELECT COUNT(*) FROM invoices").fetchone()[0],
            "monthly_summary": self.monthly_summary(),
            "tax_burden_trend": self.tax_burden_trend(months),
            "supplier_concentration": self.supplier_concentration(top_n),
            "amount_anomaly": self.amount_anomaly(z_threshold),
        }

    # ---------- 导出 ----------

    def export_excel(self, output_path: Optional[str] = None) -> Dict[str, Any]:
        """
        导出月度附件 Excel（openpyxl）
        openpyxl 不可用时自动降级为 CSV，保证功能不中断
        """
        out = Path(output_path) if output_path else \
            Path.home() / ".workbuddy" / "output" / f"ledger_report_{datetime.now():%Y%m}.xlsx"
        out.parent.mkdir(parents=True, exist_ok=True)
        report = self.full_report()

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment
        except ImportError:
            csv_path = out.with_suffix(".csv")
            rows = self.query_invoices()
            if rows:
                import csv
                with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
                    writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
                    writer.writeheader()
                    for r in rows:
                        r.pop("raw_json", None)
                        writer.writerow(r)
            return {
                "status": "degraded",
                "path": str(csv_path),
                "engine": "csv",
                "message": "未安装 openpyxl，已降级导出 CSV；执行 pip install openpyxl 可导出多页签 Excel",
            }

        wb = Workbook()
        head_font = Font(bold=True)

        def _sheet(title: str, headers: List[str], rows: List[List[Any]]):
            ws = wb.create_sheet(title)
            ws.append(headers)
            for c in ws[1]:
                c.font = head_font
                c.alignment = Alignment(horizontal="center")
            for r in rows:
                ws.append(r)
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 42)
            return ws

        wb.remove(wb.active)

        # 明细
        detail = self.query_invoices()
        if detail:
            cols = [c for c in detail[0].keys() if c != "raw_json"]
            _sheet("台账明细", cols, [[r.get(c) for c in cols] for r in detail])

        # 月度汇总
        _sheet("月度汇总", ["月份", "方向", "张数", "不含税金额", "税额", "价税合计"],
               [[r["month"], r["direction"], r["count"], r["amount"], r["tax_amount"], r["total"]]
                for r in report["monthly_summary"]])

        # 税负率趋势
        _sheet("税负率趋势",
               ["月份", "不含税销售额", "销项税额", "进项税额", "应纳税额", "税负率"],
               [[r["month"], r["output_amount"], r["output_tax"], r["input_tax"],
                 r["payable_tax"], r["tax_burden_pct"]] for r in report["tax_burden_trend"]])

        # 供应商集中度
        sc = report["supplier_concentration"]
        _sheet("供应商集中度", ["供应商", "金额", "占比"],
               [[s["supplier"], s["total"], f"{s['share'] * 100:.2f}%"] for s in sc["top_suppliers"]])
        ws = wb["供应商集中度"]
        ws.append([])
        ws.append([f"Top{sc['top_n']} 合计占比", sc["top_share_pct"]])

        # 异常波动
        an = report["amount_anomaly"]
        _sheet("异常波动", ["月份", "金额", "张数", "z-score", "状态"],
               [[s["month"], round(s["total"], 2), s["count"], s.get("z_score", 0),
                 "超限" if abs(s.get("z_score", 0)) >= an["z_threshold"] else "正常"]
                for s in an.get("series", [])])

        wb.save(str(out))
        return {"status": "success", "path": str(out), "engine": "openpyxl"}

    def close(self):
        self.conn.close()


def main():
    import argparse
    parser = argparse.ArgumentParser(description="进销项台账与月度统计")
    parser.add_argument("--db", help=f"台账数据库路径（默认 {DEFAULT_DB_PATH}）")
    parser.add_argument("--import-json", help="导入票据解析结果 JSON")
    parser.add_argument("--direction", choices=list(DIRECTIONS), default="input",
                        help="进项 input / 销项 output")
    parser.add_argument("--fill-deadlines", action="store_true", help="补算缺失的认证期限")
    parser.add_argument("--set-deadline", nargs=2, metavar=("NUMBER", "DATE"),
                        help="手工补录认证期限：--set-deadline 发票号码 YYYY-MM-DD")
    parser.add_argument("--report", action="store_true", help="输出完整报表 JSON")
    parser.add_argument("--export-excel", nargs="?", const="", help="导出 Excel 月度附件")
    parser.add_argument("--output", help="报表 JSON 输出路径")
    args = parser.parse_args()

    db = LedgerDB(args.db)

    if args.import_json:
        n = db.load_from_json(args.import_json, args.direction)
        print(f"导入完成：{n} 张")

    if args.fill_deadlines:
        n = db.fill_missing_deadlines()
        print(f"补算认证期限：{n} 张")

    if args.set_deadline:
        ok = db.set_certification_deadline(args.set_deadline[0], args.set_deadline[1])
        print("补录成功" if ok else "未找到该发票号码，请确认")

    if args.report or args.output:
        report = db.full_report()
        text = json.dumps(report, ensure_ascii=False, indent=2)
        if args.output:
            Path(args.output).write_text(text, encoding="utf-8")
            print(f"报表已保存：{args.output}")
        else:
            print(text)

    if args.export_excel is not None:
        res = db.export_excel(args.export_excel or None)
        print(json.dumps(res, ensure_ascii=False, indent=2))

    if not any([args.import_json, args.fill_deadlines, args.set_deadline,
                args.report, args.output, args.export_excel is not None]):
        parser.print_help()

    db.close()


if __name__ == "__main__":
    main()
